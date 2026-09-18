
from pathlib import Path
import shutil
import subprocess

from pypdf import PdfReader, PdfWriter
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from PIL import Image


def _check_pdf(path):
    if not Path(path).is_file():
        raise FileNotFoundError(path)
    return Path(path)


def merge_pdfs(inputs, output):
    writer = PdfWriter()
    for path in inputs:
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError(f"Encrypted PDF needs to be unlocked first: {path}")
        for page in reader.pages:
            writer.add_page(page)
    with open(output, "wb") as f:
        writer.write(f)
    return f"Merged {len(inputs)} PDF files into {output}"


def split_pdf(input_pdf, page_spec, output):
    reader = PdfReader(input_pdf)
    pages = parse_page_spec(page_spec, len(reader.pages))
    writer = PdfWriter()
    for idx in pages:
        writer.add_page(reader.pages[idx])
    with open(output, "wb") as f:
        writer.write(f)
    return f"Exported {len(pages)} page(s) to {output}"


def parse_page_spec(spec, total_pages):
    selected = set()
    for part in spec.replace(" ", "").split(","):
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            a, b = int(a), int(b)
            if a < 1 or b < a or b > total_pages:
                raise ValueError(f"Invalid range: {part}")
            selected.update(range(a - 1, b))
        else:
            n = int(part)
            if n < 1 or n > total_pages:
                raise ValueError(f"Invalid page: {part}")
            selected.add(n - 1)
    return sorted(selected)


def protect_pdf(input_pdf, output, password):
    reader = PdfReader(input_pdf)
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    # AES-256 is used when pypdf's cryptography provider is installed.
    writer.encrypt(password, algorithm="AES-256")
    with open(output, "wb") as f:
        writer.write(f)
    return f"Protected PDF saved to {output}"


def unlock_pdf(input_pdf, output, password):
    reader = PdfReader(input_pdf)
    if not reader.is_encrypted:
        raise ValueError("This PDF is not encrypted.")
    if reader.decrypt(password) == 0:
        raise ValueError("Incorrect password.")
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    with open(output, "wb") as f:
        writer.write(f)
    return f"Unlocked PDF saved to {output}"


def images_to_pdf(images, output):
    converted = []
    for src in images:
        img = Image.open(src).convert("RGB")
        converted.append(img)
    if not converted:
        raise ValueError("No images selected.")
    converted[0].save(output, save_all=True, append_images=converted[1:])
    for img in converted:
        img.close()
    return f"Created PDF from {len(images)} image(s): {output}"


def _find_libreoffice():
    candidates = [
        "soffice", "libreoffice",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for candidate in candidates:
        if shutil.which(candidate):
            return candidate
        if Path(candidate).exists():
            return candidate
    return None


def word_to_pdf(input_docx, output_pdf):
    office = _find_libreoffice()
    if office:
        out_dir = str(Path(output_pdf).resolve().parent)
        subprocess.run(
            [office, "--headless", "--convert-to", "pdf",
             "--outdir", out_dir, str(input_docx)],
            check=True, capture_output=True, text=True
        )
        generated = Path(out_dir) / (Path(input_docx).stem + ".pdf")
        if generated.resolve() != Path(output_pdf).resolve():
            shutil.move(str(generated), output_pdf)
        return f"Converted DOCX to PDF using LibreOffice: {output_pdf}"

    # Portable fallback: preserve readable text, headings and tables.
    doc = Document(input_docx)
    c = canvas.Canvas(str(output_pdf), pagesize=A4)
    width, height = A4
    x, y = 54, height - 60
    c.setFont("Helvetica", 11)

    def new_page():
        nonlocal y
        c.showPage()
        c.setFont("Helvetica", 11)
        y = height - 60

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            y -= 8
            continue
        for line in _wrap(text, 95):
            if y < 55:
                new_page()
            c.drawString(x, y, line)
            y -= 15

    for table in doc.tables:
        for row in table.rows:
            text = " | ".join(cell.text.strip() for cell in row.cells)
            for line in _wrap(text, 95):
                if y < 55:
                    new_page()
                c.drawString(x, y, line)
                y -= 15
        y -= 8

    c.save()
    return f"Converted DOCX to PDF using portable fallback: {output_pdf}"


def _wrap(text, width):
    words = text.split()
    lines, current = [], ""
    for word in words:
        if len(current) + len(word) + 1 <= width:
            current = (current + " " + word).strip()
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def pdf_to_word(input_pdf, output_docx):
    reader = PdfReader(input_pdf)
    if reader.is_encrypted:
        raise ValueError("Unlock the PDF first.")
    doc = Document()
    doc.add_heading(Path(input_pdf).stem, level=1)

    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        doc.add_heading(f"Page {page_no}", level=2)
        for line in text.splitlines():
            if line.strip():
                doc.add_paragraph(line.strip())

    doc.save(output_docx)
    return f"Converted PDF text to editable DOCX: {output_docx}"
